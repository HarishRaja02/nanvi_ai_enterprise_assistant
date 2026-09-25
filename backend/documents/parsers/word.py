from datetime import datetime
from io import BytesIO
from ..parser import DocumentParser
from ..models import Document, DocumentTable
from .common import source_for, metadata_for, clean

class WordParser(DocumentParser):
    extensions = frozenset({".docx"})
    def parse_bytes(self, *, filename: str, relative_path: str, data: bytes, modified_at: datetime | None = None) -> Document:
        try:
            from docx import Document as WordDocument
            document = WordDocument(BytesIO(data))
        except Exception as exc:
            raise ValueError("Malformed DOCX document") from exc
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        tables, table_text = [], []
        for idx, table in enumerate(document.tables, start=1):
            rows = [[clean(cell.text) for cell in row.cells] for row in table.rows]
            headers, body = (rows[0], rows[1:]) if rows else ([], [])
            tables.append(DocumentTable(headers, body, f"Table {idx}"))
            table_text.append(f"[Table {idx}]\n" + "\n".join(" | ".join(r) for r in rows))
        return Document(source_for(filename, relative_path, "docx"), metadata_for(len(data), modified_at),
                        "\n\n".join(paragraphs + table_text), tables, [f"Table {i}" for i in range(1, len(tables)+1)])
