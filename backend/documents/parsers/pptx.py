from datetime import datetime
from io import BytesIO
from ..parser import DocumentParser
from ..models import Document
from .common import source_for, metadata_for, clean

class PPTXParser(DocumentParser):
    extensions = frozenset({".pptx"})
    def parse_bytes(self, *, filename: str, relative_path: str, data: bytes, modified_at: datetime | None = None) -> Document:
        try:
            from pptx import Presentation
            presentation = Presentation(BytesIO(data))
        except Exception as exc:
            raise ValueError("Malformed PPTX document") from exc
        chunks, locations = [], []
        for index, slide in enumerate(presentation.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text = clean(shape.text)
                    if text:
                        texts.append(text)
            if texts:
                chunks.append(f"[Slide {index}]\n" + "\n".join(texts))
                locations.append(f"Slide {index}")
        return Document(source_for(filename, relative_path, "pptx"), metadata_for(len(data), modified_at),
                        "\n\n".join(chunks), [], locations)
