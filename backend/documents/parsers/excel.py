from datetime import datetime
from io import BytesIO
from ..parser import DocumentParser
from ..models import Document, DocumentTable
from .common import source_for, metadata_for, clean

class ExcelParser(DocumentParser):
    extensions = frozenset({".xlsx"})
    def parse_bytes(self, *, filename: str, relative_path: str, data: bytes, modified_at: datetime | None = None) -> Document:
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise ValueError("Malformed XLSX document") from exc
        tables, chunks, locations = [], [], []
        try:
            for sheet in workbook.worksheets:
                rows = [[clean(v) for v in row] for row in sheet.iter_rows(values_only=True)]
                rows = [r for r in rows if any(r)]
                headers, body = (rows[0], rows[1:]) if rows else ([], [])
                if len(rows) >= 2:
                    header_line = " | ".join(rows[0])
                    sep_line = " | ".join(["---"] * len(rows[0]))
                    data_lines = [" | ".join(r) for r in rows[1:]]
                    chunks.append(f"[Sheet {sheet.title}]\n" + "\n".join([header_line, sep_line, *data_lines]))
                    tables.append(DocumentTable(headers=headers, rows=body, source_label=f"Sheet {sheet.title}"))
                else:
                    chunks.append(f"[Sheet {sheet.title}]\n" + "\n".join(" | ".join(r) for r in rows))
                locations.append(f"Sheet {sheet.title}")
        finally:
            workbook.close()
        return Document(source_for(filename, relative_path, "xlsx"), metadata_for(len(data), modified_at),
                        "\n\n".join(chunks), tables, locations)
