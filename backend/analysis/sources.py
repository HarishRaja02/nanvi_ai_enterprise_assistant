from __future__ import annotations
import csv
from pathlib import Path
from typing import Any
from openpyxl import load_workbook

from backend.integrations.files.service import FileService
from backend.security.authorization import UserAttributes
from .models import DataLineage, StructuredDataset


class StructuredDataSource:
    """Loads CSV/XLSX only after FileService authorization has succeeded."""
    def __init__(self, file_service: FileService):
        self._files = file_service

    def load(self, user: UserAttributes, relative_path: str, sheet: str | None = None) -> StructuredDataset:
        data = self._files.read_file(user, relative_path)
        suffix = Path(relative_path).suffix.casefold()
        if suffix == ".csv":
            rows = list(csv.DictReader(data.decode("utf-8-sig", errors="strict").splitlines()))
            columns = tuple(rows[0].keys()) if rows else ()
        elif suffix == ".xlsx":
            wb = load_workbook(filename=__import__("io").BytesIO(data), read_only=True, data_only=True)
            ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
            values = list(ws.iter_rows(values_only=True))
            columns = tuple(str(x) for x in values[0]) if values else ()
            rows = [dict(zip(columns, row)) for row in values[1:] if any(v is not None for v in row)]
        else:
            raise ValueError("Data analysis supports CSV and XLSX files only")
        metadata = self._files.get_metadata(user, relative_path)
        top_level = relative_path.replace("\\", "/").split("/", 1)[0].casefold()
        restricted_department = {"hr": "HR", "finance": "Finance"}.get(top_level)
        lineage = DataLineage("file", relative_path, metadata.name, columns, tenant_id=user.tenant_id, department=restricted_department, restricted_department=restricted_department, resource_type="company_file")
        return StructuredDataset(columns, tuple(rows), (lineage,))
