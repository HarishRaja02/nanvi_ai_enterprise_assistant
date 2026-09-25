from __future__ import annotations

from abc import ABC, abstractmethod
from html import escape
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from docx import Document
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches, Pt as PptPt

from .models import ReportRequest


class ReportGenerator(ABC):
    extension: str

    @abstractmethod
    def generate(self, request: ReportRequest, destination: Path) -> None:
        raise NotImplementedError

    @staticmethod
    def _rows(data: Any) -> list[dict[str, Any]]:
        if isinstance(data, list) and all(isinstance(row, dict) for row in data):
            return data
        if isinstance(data, tuple) and all(isinstance(row, dict) for row in data):
            return list(data)
        if isinstance(data, dict):
            # Scalar/keyed analysis results become a compact key/value table.
            return [{"metric": key, "value": value} for key, value in data.items()]
        return [{"value": data}]

    @classmethod
    def _table_data(cls, data: Any) -> tuple[list[str], list[list[str]]]:
        rows = cls._rows(data)
        columns: list[str] = []
        for row in rows:
            for key in row:
                if key not in columns:
                    columns.append(str(key))
        values = [[cls._display(row.get(column)) for column in columns] for row in rows]
        return columns, values

    @staticmethod
    def _display(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, float):
            return f"{value:,.4f}".rstrip("0").rstrip(".")
        return str(value)

    @staticmethod
    def _excel_safe(value: Any) -> Any:
        # Prevent formula injection when untrusted strings are opened in Excel,
        # while preserving genuine numeric/date values for correct calculations.
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            return "'" + value
        return value

    @staticmethod
    def _lineage_rows(request: ReportRequest) -> list[list[str]]:
        # Reports expose human-safe lineage only; never embed raw SQL or source IDs.
        return [
            [lineage.source_type, lineage.source_label,
             ", ".join(lineage.columns), "; ".join(lineage.filters)]
            for lineage in request.lineage
        ]


class ExcelReportGenerator(ReportGenerator):
    extension = "xlsx"

    def generate(self, request: ReportRequest, destination: Path) -> None:
        from openpyxl.styles import PatternFill, Alignment, Border, Side
        from datetime import datetime, timezone

        wb = Workbook()
        ws = wb.active
        ws.title = "Report"

        # Document Header
        ws.append([self._excel_safe(request.title)])
        ws["A1"].font = Font(name="Calibri", bold=True, size=16, color="0F172A")

        if request.description:
            ws.append([self._excel_safe(request.description)])
            ws["A2"].font = Font(name="Calibri", italic=True, size=10, color="334155")
        ws.append([])

        rows_raw = self._rows(request.data)
        columns: list[str] = []
        for row in rows_raw:
            for key in row:
                if key not in columns:
                    columns.append(str(key))

        ws.append(columns or ["No data"])
        header_row = ws.max_row

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        for cell in ws[header_row]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", horizontal="left")

        for r_idx, row in enumerate(rows_raw):
            row_vals = []
            for c_idx, column in enumerate(columns):
                val = row.get(column)
                row_vals.append(self._excel_safe(val))
            ws.append(row_vals)
            curr_row = ws.max_row
            is_even = r_idx % 2 == 1
            for cell in ws[curr_row]:
                cell.font = Font(name="Calibri", size=10, color="0F172A")
                cell.border = thin_border
                if is_even:
                    cell.fill = zebra_fill

        for idx, column in enumerate(columns, 1):
            col_letter = get_column_letter(idx)
            max_len = max(len(str(column)), 12)
            for r in range(header_row, ws.max_row + 1):
                cell_val = ws[f"{col_letter}{r}"].value
                if cell_val is not None:
                    max_len = max(max_len, len(str(cell_val)))
            ws.column_dimensions[col_letter].width = min(max_len + 3, 40)

        lineage = wb.create_sheet("Sources")
        lineage.append(["Source Type", "Source Label", "Columns", "Filters"])
        lineage["A1"].font = Font(bold=True)
        lineage["B1"].font = Font(bold=True)
        lineage["C1"].font = Font(bold=True)
        lineage["D1"].font = Font(bold=True)
        for row in self._lineage_rows(request):
            lineage.append([self._excel_safe(str(value)) for value in row])
        for col_idx in range(1, 5):
            lineage.column_dimensions[get_column_letter(col_idx)].width = 25

        wb.save(destination)


class PDFReportGenerator(ReportGenerator):
    extension = "pdf"

    def generate(self, request: ReportRequest, destination: Path) -> None:
        from datetime import datetime, timezone

        columns, rows = self._table_data(request.data)
        if not columns:
            columns, rows = ["Result"], [["No data available"]]

        display_cols = [str(c).title().replace("_", " ") for c in columns]

        doc = SimpleDocTemplate(
            str(destination),
            pagesize=landscape(A4),
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36,
        )
        styles = getSampleStyleSheet()

        title_style = styles["Title"]
        title_style.fontName = "Helvetica-Bold"
        title_style.fontSize = 20
        title_style.leading = 24
        title_style.textColor = colors.HexColor("#0F172A")
        title_style.alignment = 0

        body_style = styles["BodyText"]
        body_style.fontName = "Helvetica"
        body_style.fontSize = 10
        body_style.leading = 14
        body_style.textColor = colors.HexColor("#334155")

        meta_style = styles["Italic"]
        meta_style.fontName = "Helvetica-Oblique"
        meta_style.fontSize = 8.5
        meta_style.leading = 11
        meta_style.textColor = colors.HexColor("#64748B")

        h2_style = styles["Heading2"]
        h2_style.fontName = "Helvetica-Bold"
        h2_style.fontSize = 12
        h2_style.leading = 16
        h2_style.textColor = colors.HexColor("#0F172A")

        now_str = datetime.now(timezone.utc).strftime("%B %d, %Y %H:%M UTC")
        story = [
            Paragraph(escape(request.title), title_style),
            Spacer(1, 4),
            Paragraph(escape(f"CONFIDENTIAL ENTERPRISE DOCUMENT  •  Generated: {now_str}  •  Tenant: {request.tenant_id}"), meta_style),
            Spacer(1, 14),
        ]

        if request.description:
            story.append(Paragraph(escape(request.description), body_style))
            story.append(Spacer(1, 12))

        formatted_rows = [[str(val) for val in row] for row in rows]
        table_data = [display_cols] + formatted_rows
        table = Table(table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 9),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
            ("TOPPADDING", (0, 0), (-1, 0), 6),
            ("TEXTCOLOR", (0, 1), (-1, -1), colors.HexColor("#0F172A")),
            ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 1), (-1, -1), 8.5),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BOTTOMPADDING", (0, 1), (-1, -1), 5),
            ("TOPPADDING", (0, 1), (-1, -1), 5),
        ]))
        story.append(table)
        story.append(Spacer(1, 14))

        story.append(Paragraph("Governance & Audit Lineage", h2_style))
        for lineage in request.lineage:
            story.append(Paragraph(
                escape(f"• {lineage.source_type.upper()}: {lineage.source_label} (Columns: {', '.join(lineage.columns)})"),
                meta_style,
            ))

        doc.build(story)


class WordReportGenerator(ReportGenerator):
    extension = "docx"

    def generate(self, request: ReportRequest, destination: Path) -> None:
        from datetime import datetime, timezone
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls

        document = Document()

        # Page setup: Standard Letter with 1-inch margins
        for section in document.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        # 1. Document Header / Title (20pt Bold, Deep Charcoal / Black Ink)
        title_para = document.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        title_run = title_para.add_run(request.title[:200])
        title_run.font.name = "Calibri"
        title_run.font.size = Pt(20)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(15, 23, 42)
        title_para.paragraph_format.space_after = Pt(4)
        title_para.paragraph_format.space_before = Pt(0)

        # 2. Metadata Banner
        meta_para = document.add_paragraph()
        meta_run1 = meta_para.add_run("CONFIDENTIAL ENTERPRISE DOCUMENT  •  ")
        meta_run1.font.name = "Calibri"
        meta_run1.font.size = Pt(9)
        meta_run1.font.bold = True
        meta_run1.font.color.rgb = RGBColor(37, 99, 235)

        now_str = datetime.now(timezone.utc).strftime("%B %d, %Y %H:%M UTC")
        meta_run2 = meta_para.add_run(f"Generated: {now_str}  •  Tenant: {request.tenant_id}")
        meta_run2.font.name = "Calibri"
        meta_run2.font.size = Pt(9)
        meta_run2.font.color.rgb = RGBColor(100, 116, 139)
        meta_para.paragraph_format.space_after = Pt(14)

        # Divider line
        div_para = document.add_paragraph()
        div_run = div_para.add_run("―" * 54)
        div_run.font.size = Pt(8)
        div_run.font.color.rgb = RGBColor(226, 232, 240)
        div_para.paragraph_format.space_after = Pt(12)

        # 3. Executive Description / Overview
        if request.description:
            desc_heading = document.add_paragraph()
            desc_h_run = desc_heading.add_run("Executive Overview")
            desc_h_run.font.name = "Calibri"
            desc_h_run.font.size = Pt(13)
            desc_h_run.font.bold = True
            desc_h_run.font.color.rgb = RGBColor(15, 23, 42)
            desc_heading.paragraph_format.space_after = Pt(4)

            desc_para = document.add_paragraph()
            desc_run = desc_para.add_run(request.description[:4000])
            desc_run.font.name = "Calibri"
            desc_run.font.size = Pt(10.5)
            desc_run.font.color.rgb = RGBColor(51, 65, 85)
            desc_para.paragraph_format.space_after = Pt(14)

        # 4. Data Section Heading
        data_heading = document.add_paragraph()
        data_h_run = data_heading.add_run("Authorized Record Schedule")
        data_h_run.font.name = "Calibri"
        data_h_run.font.size = Pt(13)
        data_h_run.font.bold = True
        data_h_run.font.color.rgb = RGBColor(15, 23, 42)
        data_heading.paragraph_format.space_after = Pt(6)

        # 5. Formatted Professional Table
        columns, rows = self._table_data(request.data)
        if not columns:
            columns, rows = ["Result"], [["No data available"]]

        num_cols = len(columns)
        table = document.add_table(rows=1, cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        header_row = table.rows[0]
        header_tr_pr = header_row._tr.get_or_add_trPr()
        header_tr_pr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for idx, (cell, val) in enumerate(zip(header_row.cells, columns)):
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1E293B"/>')
            cell._tc.get_or_add_tcPr().append(shd)

            tc_mar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="120" w:type="dxa"/><w:bottom w:w="120" w:type="dxa"/><w:left w:w="140" w:type="dxa"/><w:right w:w="140" w:type="dxa"/></w:tcMar>')
            cell._tc.get_or_add_tcPr().append(tc_mar)

            cell.text = str(val).title().replace("_", " ")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.bold = True
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(255, 255, 255)

        for r_idx, row in enumerate(rows):
            new_row = table.add_row()
            bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"

            for c_idx, (cell, val) in enumerate(zip(new_row.cells, row)):
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
                cell._tc.get_or_add_tcPr().append(shd)

                tc_mar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="90" w:type="dxa"/><w:bottom w:w="90" w:type="dxa"/><w:left w:w="140" w:type="dxa"/><w:right w:w="140" w:type="dxa"/></w:tcMar>')
                cell._tc.get_or_add_tcPr().append(tc_mar)

                disp_val = str(val)
                cell.text = disp_val
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if disp_val.startswith("$") else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(15, 23, 42)

        tbl_pr = table._tbl.tblPr
        tbl_borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
            f'<w:insideV w:val="none"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tbl_pr.append(tbl_borders)

        # 6. Sources & Compliance Footer
        document.add_paragraph().paragraph_format.space_before = Pt(16)
        src_h = document.add_paragraph()
        src_h_run = src_h.add_run("Governance & Audit Lineage")
        src_h_run.font.name = "Calibri"
        src_h_run.font.size = Pt(11)
        src_h_run.font.bold = True
        src_h_run.font.color.rgb = RGBColor(71, 85, 105)
        src_h.paragraph_format.space_after = Pt(4)

        for lineage in request.lineage:
            src_p = document.add_paragraph()
            src_p.paragraph_format.space_after = Pt(2)
            b_run = src_p.add_run(f"• {lineage.source_type.upper()}: ")
            b_run.font.name = "Calibri"
            b_run.font.bold = True
            b_run.font.size = Pt(9)
            b_run.font.color.rgb = RGBColor(100, 116, 139)

            d_run = src_p.add_run(f"{lineage.source_label} (Columns: {', '.join(lineage.columns)})")
            d_run.font.name = "Calibri"
            d_run.font.size = Pt(9)
            d_run.font.color.rgb = RGBColor(100, 116, 139)

        document.save(destination)


class TextReportGenerator(ReportGenerator):
    extension = "txt"

    def generate(self, request: ReportRequest, destination: Path) -> None:
        from datetime import datetime, timezone

        columns, rows = self._table_data(request.data)
        if not columns:
            columns, rows = ["Result"], [["No data available"]]

        display_cols = [str(c).title().replace("_", " ") for c in columns]
        col_widths = [len(c) for c in display_cols]
        for row in rows:
            for idx, val in enumerate(row):
                if idx < len(col_widths):
                    col_widths[idx] = max(col_widths[idx], len(str(val)))

        lines = [
            "=" * 80,
            f"  {request.title.upper()}",
            "=" * 80,
            f"Classification: CONFIDENTIAL ENTERPRISE DOCUMENT",
            f"Generated:      {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}",
            f"Tenant ID:      {request.tenant_id}",
            f"Format:         Plain Text (.txt)",
            "-" * 80,
        ]
        if request.description:
            lines.extend([request.description, "-" * 80])
        lines.append("")

        header_str = " | ".join(c.ljust(col_widths[i]) for i, c in enumerate(display_cols))
        sep_str = "-+-".join("-" * col_widths[i] for i in range(len(col_widths)))
        lines.append(header_str)
        lines.append(sep_str)

        for row in rows:
            row_str = " | ".join(str(row[i] if i < len(row) else "").ljust(col_widths[i]) for i in range(len(col_widths)))
            lines.append(row_str)

        lines.extend([
            "",
            "-" * 80,
            "AUDIT & DATA LINEAGE:",
        ])
        for lineage in request.lineage:
            lines.append(f"• {lineage.source_type.upper()}: {lineage.source_label} (Columns: {', '.join(lineage.columns)})")
        lines.extend([
            "=" * 80,
            "Nanvi AI Enterprise Assistant — Official Document Output",
            "=" * 80,
        ])
        destination.write_text("\n".join(lines), encoding="utf-8")


class PowerPointReportGenerator(ReportGenerator):
    extension = "pptx"

    def generate(self, request: ReportRequest, destination: Path) -> None:
        prs = Presentation()
        title_slide = prs.slides.add_slide(prs.slide_layouts[0])
        title_slide.shapes.title.text = request.title[:200]
        subtitle = title_slide.placeholders[1]
        subtitle.text = (request.description or "Nanvi report")[:1000]

        columns, rows = self._table_data(request.data)
        if not columns:
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = "Data"
            box = slide.shapes.add_textbox(Inches(0.7), Inches(1.5), Inches(11.8), Inches(2.0))
            box.text_frame.text = "No data available"
        else:
            chunk_size = 24
            for offset in range(0, len(rows), chunk_size):
                chunk = rows[offset:offset + chunk_size]
                slide = prs.slides.add_slide(prs.slide_layouts[5])
                slide.shapes.title.text = f"Data{(' — rows ' + str(offset + 1) + '-' + str(offset + len(chunk))) if len(rows) > chunk_size else ''}"
                table = slide.shapes.add_table(len(chunk) + 1, len(columns), Inches(0.3), Inches(1.2), Inches(12.7), Inches(5.8)).table
                for c, header in enumerate(columns):
                    table.cell(0, c).text = str(header)[:80]
                for r, row in enumerate(chunk, start=1):
                    for c, value in enumerate(row):
                        table.cell(r, c).text = str(value)[:200]
                font_size = max(7, min(11, int(90 / max(len(columns), 1))))
                for row_cells in table.rows:
                    for cell in row_cells.cells:
                        for paragraph in cell.text_frame.paragraphs:
                            for run in paragraph.runs:
                                run.font.size = PptPt(font_size)

        source_slide = prs.slides.add_slide(prs.slide_layouts[5])
        source_slide.shapes.title.text = "Sources / Lineage"
        source_lines = self._lineage_rows(request)
        text = "\n".join(
            f"{source_type}: {label} | columns: {columns} | filters: {filters or 'none'}"
            for source_type, label, columns, filters in source_lines
        ) or "No source lineage supplied"
        box = source_slide.shapes.add_textbox(Inches(0.7), Inches(1.4), Inches(11.7), Inches(5.5))
        box.text_frame.text = text[:10000]
        prs.save(destination)
