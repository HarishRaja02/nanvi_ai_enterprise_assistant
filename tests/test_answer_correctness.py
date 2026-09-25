from __future__ import annotations
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from backend.analysis import (
    AnalysisOperation, AnalysisRequest, DataLineage, DeterministicAnalysisEngine,
    StructuredDataset,
)
from backend.documents import CSVParser, ExcelParser, PDFParser, WordParser
from backend.integrations.database.models import QueryRequest, QueryResult, TablePolicy
from backend.integrations.database.sql_validation import ReadOnlySQLValidator, SQLValidationPipeline
from backend.integrations.database.service import DatabaseService
from backend.integrations.email.models import EmailAddress, EmailMessage, EmailProviderContext, EmailSearchRequest
from backend.integrations.email.repository import EmailRepository
from backend.integrations.email.service import EmailService
from backend.retrieval import (
    AccessControlMetadata, DocumentChunker, KeywordRetriever, KnowledgeRetrievalService,
    NoOpReranker, SourceMetadata, RetrievalRequest,
)
from backend.security.audit import AuditLogger, InMemoryAuditSink
from backend.security.authorization import AuthorizationService, UserAttributes
from backend.security.authorization.rbac import Role

USER = UserAttributes("gold-user", "gold-tenant", "Projects", frozenset({Role.EMPLOYEE}))

class GoldenDB:
    def __init__(self):
        self.rows = (
            ("A", "2026-01-10", 10000),
            ("A", "2026-02-10", 15000),
            ("B", "2026-01-12", 20000),
            ("B", "2026-02-20", 30000),
            ("C", "2026-02-25", 5000),
        )
        self.calls = []
    def execute_read(self, request):
        self.calls.append(request)
        return QueryResult(("customer", "date", "sales"), self.rows)

class GoldenEmail(EmailRepository):
    def __init__(self):
        self.messages = (
            EmailMessage("e1", "c1", "Project Atlas renewal", EmailAddress("alice@example.test", "Alice"), (), (),
                         datetime(2026,2,5,tzinfo=timezone.utc), None, "Atlas renewal approved for 30,000.", False),
            EmailMessage("e2", "c2", "Project Beta update", EmailAddress("bob@example.test", "Bob"), (), (),
                         datetime(2026,1,15,tzinfo=timezone.utc), None, "Beta is delayed by two weeks.", False),
        )
    def search(self, context, request):
        items = list(self.messages)
        if request.query:
            q=request.query.casefold(); items=[m for m in items if q in (m.subject+" "+(m.body_preview or "")).casefold()]
        if request.after: items=[m for m in items if m.received_at and m.received_at >= request.after]
        if request.before: items=[m for m in items if m.received_at and m.received_at < request.before]
        return items[:request.max_results]
    def get_thread(self, context, conversation_id): raise NotImplementedError
    def get_message(self, context, message_id): raise NotImplementedError


def build_files(root: Path):
    root.mkdir(parents=True, exist_ok=True)
    (root/"facts.txt").write_text("The Atlas project owner is Priya. The project deadline is 2026-03-31.", encoding="utf-8")
    (root/"notes.csv").write_text("customer,date,sales\nA,2026-01-10,10000\nA,2026-02-10,15000\nB,2026-01-12,20000\nB,2026-02-20,30000\nC,2026-02-25,5000\n", encoding="utf-8")
    wb=Workbook(); ws=wb.active; ws.title="Sales"; ws.append(["customer","date","sales"]); [ws.append(list(r)) for r in [("A","2026-01-10",10000),("A","2026-02-10",15000),("B","2026-01-12",20000),("B","2026-02-20",30000),("C","2026-02-25",5000)]]; wb.save(root/"sales.xlsx"); wb.close()
    d=Document(); d.add_paragraph("Atlas Word document: owner is Priya and deadline is 2026-03-31."); d.save(root/"atlas.docx")
    c=canvas.Canvas(str(root/"atlas.pdf")); c.drawString(72,720,"Atlas PDF: owner is Priya and deadline is 2026-03-31."); c.save()


def retrieval_service(root: Path):
    auth=AuthorizationService(); audit=AuditLogger(InMemoryAuditSink())
    r=KnowledgeRetrievalService(None,None,KeywordRetriever(),NoOpReranker(),auth,audit,DocumentChunker())
    parsers={".pdf":PDFParser(),".docx":WordParser(),".xlsx":ExcelParser(),".csv":CSVParser()}
    for p in root.iterdir():
        if p.suffix == ".txt": text=p.read_text(encoding="utf-8")
        else: text=parsers[p.suffix].parse_bytes(filename=p.name, relative_path=p.name, data=p.read_bytes()).text
        r.index(text, SourceMetadata(p.suffix.lstrip("."), p.name, p.name, p.name, access=AccessControlMetadata("gold-tenant")))
    return r


def dataset():
    rows=(
        {"customer":"A","date":"2026-01-10","sales":10000},
        {"customer":"A","date":"2026-02-10","sales":15000},
        {"customer":"B","date":"2026-01-12","sales":20000},
        {"customer":"B","date":"2026-02-20","sales":30000},
        {"customer":"C","date":"2026-02-25","sales":5000},
    )
    lineage=(DataLineage("csv","notes.csv","Sales CSV",("customer","date","sales"), tenant_id="gold-tenant", resource_type="file"),)
    return StructuredDataset(("customer","date","sales"),rows,lineage)


def test_golden_document_answers(tmp_path):
    build_files(tmp_path); r=retrieval_service(tmp_path)
    cases=[
        ("Who owns the Atlas project?", "Priya", "facts.txt"),
        ("What is the Atlas deadline?", "2026-03-31", "facts.txt"),
        ("What does the PDF say about the Atlas owner?", "Priya", "atlas.pdf"),
        ("What does the Word document say about the deadline?", "2026-03-31", "atlas.docx"),
        ("What customer sales are in the Excel file?", "A", "sales.xlsx"),
        ("What customer sales are in the CSV?", "B", "notes.csv"),
    ]
    for q, expected, source in cases:
        result=r.retrieve(USER, RetrievalRequest(q, top_k=3))
        text="\n".join(h.chunk.content for h in result.hits)
        assert expected in text, (q, expected, text)
        assert any(h.chunk.source.source_id == source for h in result.hits), (q, source, [h.chunk.source.source_id for h in result.hits])


def test_golden_sql_and_analysis_answers():
    auth=AuthorizationService(); audit=AuditLogger(InMemoryAuditSink()); repo=GoldenDB()
    validator=SQLValidationPipeline(ReadOnlySQLValidator(TablePolicy(frozenset({("public","sales")}), frozenset({("public","sales","customer"), ("public","sales","sales"), ("public","sales","date")})), require_column_policy=True))
    db=DatabaseService(repo,auth,audit,validator,"gold-tenant")
    result=db.execute_read(USER, QueryRequest("SELECT customer, date, sales FROM sales LIMIT 100"))
    assert len(result.rows)==5
    ds=dataset(); engine=DeterministicAnalysisEngine()
    total=engine.execute(AnalysisRequest(AnalysisOperation.TOTAL,ds,column="sales"))
    avg=engine.execute(AnalysisRequest(AnalysisOperation.AVERAGE,ds,column="sales"))
    feb=engine.execute(AnalysisRequest(AnalysisOperation.TOTAL,ds,column="sales",filter_column="date",filter_operator=">=",filter_value="2026-02-01"))
    grouped=engine.execute(AnalysisRequest(AnalysisOperation.GROUP,ds,column="sales",group_by="customer"))
    assert total.value == 80000
    assert avg.value == 16000
    assert feb.value == 50000
    assert grouped.value["B"]["total"] == 50000
    assert grouped.value["B"]["average"] == 25000


def test_golden_email_answers():
    auth=AuthorizationService(); audit=AuditLogger(InMemoryAuditSink()); email=EmailService(GoldenEmail(),auth,audit)
    ctx=EmailProviderContext("gold-user","gold-tenant","test-token",frozenset({"Mail.Read"}))
    all_mail=email.search(USER,ctx,EmailSearchRequest(query="Atlas"))
    assert len(all_mail)==1 and all_mail[0].body_preview == "Atlas renewal approved for 30,000."
    feb=email.search(USER,ctx,EmailSearchRequest(after=datetime(2026,2,1,tzinfo=timezone.utc)))
    assert len(feb)==1 and feb[0].id=="e1"


def test_golden_missing_data_is_explicit(tmp_path):
    build_files(tmp_path); r=retrieval_service(tmp_path)
    result=r.retrieve(USER, RetrievalRequest("What is the Atlas budget?", top_k=3))
    text="\n".join(h.chunk.content for h in result.hits)
    assert "budget" not in text.casefold()
