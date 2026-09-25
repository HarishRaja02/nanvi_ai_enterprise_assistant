from pathlib import Path
import pytest
from backend.documents import DocumentProcessingService, PDFParser, WordParser, ExcelParser, CSVParser, PPTXParser, TextParser, UnsupportedDocumentTypeError

def service():
    return DocumentProcessingService([PDFParser(), WordParser(), ExcelParser(), CSVParser(), PPTXParser(), TextParser()])

def parse_file(svc, p):
    return svc.parse_bytes(filename=p.name, relative_path=p.name, data=p.read_bytes())

def test_csv_parser(tmp_path: Path):
    p = tmp_path / "customers.csv"; p.write_text("Name,Department\nAlice,Projects\nBob,Finance\n", encoding="utf-8")
    doc = parse_file(service(), p)
    assert doc.filename == "customers.csv" and doc.source.file_type == "csv" and "Alice" in doc.text
    assert doc.tables[0].headers == ["Name", "Department"]

def test_docx_parser(tmp_path: Path):
    from docx import Document
    p = tmp_path / "report.docx"; d = Document(); d.add_paragraph("Quarterly project report")
    t = d.add_table(rows=2, cols=2); t.cell(0,0).text="Name"; t.cell(0,1).text="Status"; t.cell(1,0).text="Project A"; t.cell(1,1).text="Active"; d.save(p)
    doc = parse_file(service(), p); assert "Quarterly project report" in doc.text; assert doc.tables[0].rows == [["Project A", "Active"]]

def test_xlsx_parser(tmp_path: Path):
    from openpyxl import Workbook
    p = tmp_path / "finance.xlsx"; wb=Workbook(); ws=wb.active; ws.title="Finance"; ws.append(["Employee","Amount"]); ws.append(["Alice",1000]); wb.save(p); wb.close()
    doc=parse_file(service(),p); assert "[Sheet Finance]" in doc.text; assert doc.tables[0].rows[0]==["Alice","1000"]; assert doc.locations==["Sheet Finance"]

def test_pdf_parser(tmp_path: Path):
    from reportlab.pdfgen import canvas
    p=tmp_path/"brief.pdf"; c=canvas.Canvas(str(p)); c.drawString(72,720,"PDF source text"); c.save()
    doc=parse_file(service(),p); assert "PDF source text" in doc.text; assert doc.locations==["Page 1"]

def test_pptx_parser(tmp_path: Path):
    from pptx import Presentation
    p=tmp_path/"brief.pptx"; prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[5]); slide.shapes.title.text="Quarterly Review"; prs.save(p)
    doc=parse_file(service(),p); assert "Quarterly Review" in doc.text; assert doc.locations==["Slide 1"]

def test_txt_parser(tmp_path: Path):
    p=tmp_path/"notes.txt"; p.write_text("Project Atlas is active.",encoding="utf-8")
    doc=parse_file(service(),p); assert doc.text=="Project Atlas is active."; assert doc.source.file_type=="txt"

def test_unsupported(tmp_path: Path):
    p=tmp_path/"data.exe"; p.write_bytes(b"MZ")
    with pytest.raises(UnsupportedDocumentTypeError): service().parse_bytes(filename=p.name, relative_path=p.name, data=p.read_bytes())

def test_source_metadata_retained_without_absolute_path(tmp_path: Path):
    p=tmp_path/"source.csv"; p.write_text("A,B\n1,2\n")
    doc=parse_file(service(),p); assert doc.source.path=="source.csv"; assert not str(tmp_path) in doc.source.path; assert doc.metadata.size_bytes==p.stat().st_size
