from __future__ import annotations
import csv
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document as WordDocument
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from backend.documents import CSVParser, DocumentProcessingService, ExcelParser, PDFParser, WordParser
from backend.integrations.files import LocalFileRepository, FileService
from backend.retrieval import (
    AccessControlMetadata, EmbeddingProvider, InMemoryVectorStore, KeywordRetriever,
    KnowledgeRetrievalService, RetrievalRequest, Reranker, SourceMetadata,
)
from backend.security.audit import AuditLogger, InMemoryAuditSink
from backend.security.authorization import AuthorizationService, UserAttributes
from backend.retrieval.ingestion import source_from_email
from backend.integrations.email.models import EmailMessage


class EvalEmbedding(EmbeddingProvider):
    """Deterministic test embedding: topic buckets only, never used in production."""
    TOPICS = ("revenue", "salary", "project", "renewal", "policy", "inventory")

    def embed(self, text: str) -> list[float]:
        t = text.lower()
        return [float(k in t) for k in self.TOPICS]


class StableReranker(Reranker):
    def rerank(self, query, hits, top_k):
        return hits[:top_k]


def build_service():
    sink = InMemoryAuditSink()
    svc = KnowledgeRetrievalService(
        EvalEmbedding(), InMemoryVectorStore(), KeywordRetriever(), StableReranker(),
        AuthorizationService(), AuditLogger(sink)
    )
    return svc, sink


def src(source_id, filename, source_type="local_file", *, tenant="eval", owner=None, department=None, restricted=None, sheet=None, page=None):
    return SourceMetadata(
        source_type=source_type, source_id=source_id, filename=filename, path=source_id,
        sheet=sheet, page=page,
        access=AccessControlMetadata(
            tenant_id=tenant, owner_id=owner, department=department,
            restricted_department=restricted, resource_type=("email_mailbox" if source_type == "email" else "knowledge_source"),
        ),
    )


def user(uid="u1", dept="Projects", roles=frozenset({"Employee"}), tenant="eval"):
    return UserAttributes(uid, tenant, dept, roles)


def test_controlled_retrieval_evaluation_dataset():
    svc, _ = build_service()
    docs = [
        ("projects/roadmap.txt", "Project Atlas launches on 2026-10-15. Milestone review is 2026-10-01."),
        ("finance/revenue.csv", "Quarter,Revenue\nQ1,120000\nQ2,150000\nQ3,130000"),
        ("hr/salary.docx", "Employee compensation is confidential. Alice salary is 90000."),
        ("customers/renewal.pdf", "Customer renewal for Acme is due on 2026-11-30."),
        ("inventory.xlsx", "SKU | Stock\nA-100 | 42\nB-200 | 17"),
        ("policy.txt", "Remote work policy allows three remote days per week."),
    ]
    for sid, text in docs:
        restricted = "Finance" if sid.startswith("finance/") else ("HR" if sid.startswith("hr/") else None)
        svc.index(text, src(sid, Path(sid).name, restricted=restricted, department=restricted))

    email = EmailMessage("mail-1", "thread-1", "Acme renewal", None, (), (), None, None, "Acme renewal is due on 2026-11-30.", False)
    email_source = source_from_email(email, "eval", "u1")
    svc.index(email.body_preview, email_source)

    cases = [
        ("Project Atlas launch date", "projects/roadmap.txt", "2026-10-15"),
        ("Q2 revenue", "finance/revenue.csv", "150000"),
        ("employee salary Alice", None, "NO_ACCESS"),
        ("Acme renewal due", "customers/renewal.pdf", "2026-11-30"),
        ("stock A-100", "inventory.xlsx", "42"),
        ("remote work days", "policy.txt", "three"),
        ("Acme renewal email", "mail-1", "2026-11-30"),
        ("something not in corpus", None, "EMPTY"),
    ]

    results = []
    for question, expected_source, expected_answer in cases:
        query_user = user()
        if expected_source == "finance/revenue.csv":
            query_user = user(dept="Finance", roles=frozenset({"Finance"}))
        elif expected_source == "mail-1":
            query_user = user()
        result = svc.retrieve(query_user, RetrievalRequest(question, top_k=3, candidate_k=10))
        if expected_source is None:
            correct = len(result.hits) == 0
            actual = "NO_ACCESS" if correct else result.hits[0].chunk.content
            source = None if correct else result.hits[0].chunk.source.source_id
        elif expected_answer == "EMPTY":
            correct = len(result.hits) == 0
            actual = "EMPTY" if correct else result.hits[0].chunk.content
            source = None if correct else result.hits[0].chunk.source.source_id
        else:
            matching = [h for h in result.hits if h.chunk.source.source_id == expected_source and expected_answer.lower() in h.chunk.content.lower()]
            correct = bool(matching)
            actual = matching[0].chunk.content if matching else (result.hits[0].chunk.content if result.hits else "EMPTY")
            source = result.hits[0].chunk.source.source_id if result.hits else None
        results.append((question, expected_answer, actual, source, correct))

    assert all(row[-1] for row in results), results

    # Source preservation: every returned hit retains source identity.
    for question, expected_answer, _, _, _ in results:
        query_user = user(dept="Finance", roles=frozenset({"Finance"})) if question == "Q2 revenue" else user()
        rr = svc.retrieve(query_user, RetrievalRequest(question, top_k=3, candidate_k=10))
        assert all(h.chunk.source.source_id and h.chunk.source.filename for h in rr.hits)


def test_all_supported_parsers_feed_retrieval():
    with TemporaryDirectory() as td:
        root = Path(td)
        txt = root / "notes.txt"
        txt.write_text("Atlas budget is 250000 and owner is Priya.", encoding="utf-8")

        csv_path = root / "sales.csv"
        with csv_path.open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f); w.writerow(["Month", "Revenue"]); w.writerow(["Jan", "100"]); w.writerow(["Feb", "200"])

        xlsx = root / "inventory.xlsx"
        wb = Workbook(); ws = wb.active; ws.title = "Inventory"; ws.append(["SKU", "Stock"]); ws.append(["A-1", 42]); wb.save(xlsx)

        docx = root / "people.docx"
        wd = WordDocument(); wd.add_paragraph("Alice owns the Atlas project."); wd.save(docx)

        pdf = root / "renewal.pdf"
        c = canvas.Canvas(str(pdf)); c.drawString(72, 720, "Acme renewal date is 2026-11-30."); c.save()

        service = DocumentProcessingService([PDFParser(), WordParser(), ExcelParser(), CSVParser()])
        docs = [service.parse_bytes(filename=p.name, relative_path=p.name, data=p.read_bytes()) for p in (pdf, docx, xlsx, csv_path)]
        assert "Acme renewal date" in docs[0].text
        assert "Alice owns" in docs[1].text
        assert "A-1" in docs[2].text
        assert "Jan | 100" in docs[3].text
        assert all(d.source.filename for d in docs)
        assert docs[0].locations == ["Page 1"]
        assert docs[2].locations == ["Sheet Inventory"]

        # TXT is directly indexed as a supported local file source.
        svc, _ = build_service()
        svc.index(txt.read_text(), src(str(txt), txt.name))
        result = svc.retrieve(user(), RetrievalRequest("Atlas budget", top_k=1, candidate_k=5))
        assert result.hits[0].chunk.source.filename == "notes.txt"


def test_permission_filtering_prevents_hr_and_cross_tenant_exposure():
    svc, sink = build_service()
    svc.index("Alice salary is 90000.", src("hr/alice.docx", "alice.docx", restricted="HR", department="HR"))
    svc.index("Other tenant salary is 77777.", src("finance/other.csv", "other.csv", tenant="other", restricted="Finance", department="Finance"))
    svc.index("Project salary budget is 5000.", src("projects/budget.txt", "budget.txt"))
    result = svc.retrieve(user(), RetrievalRequest("salary", top_k=5, candidate_k=10))
    assert all("90000" not in h.chunk.content and "77777" not in h.chunk.content for h in result.hits)
    assert sink.events[-1].metadata["denied_candidates"] >= 2


def test_duplicate_chunks_are_removed_from_final_results():
    svc, _ = build_service()
    text = "Project Atlas status is active."
    svc.index(text, src("projects/status.txt", "status.txt"))
    svc.index(text, src("projects/status-copy.txt", "status-copy.txt"))
    result = svc.retrieve(user(), RetrievalRequest("Project Atlas status", top_k=10, candidate_k=20))
    ids = [h.chunk.chunk_id for h in result.hits]
    assert len(ids) == len(set(ids))


def test_zero_similarity_vector_items_do_not_pollute_results():
    svc, _ = build_service()
    svc.index("Project Atlas roadmap.", src("projects/roadmap.txt", "roadmap.txt"))
    svc.index("Revenue Q2 is 150000.", src("finance/revenue.csv", "revenue.csv", restricted="Finance", department="Finance"))
    result = svc.retrieve(user(), RetrievalRequest("question absent from corpus", top_k=5, candidate_k=5))
    assert result.hits == ()


def test_authorized_relevant_result_survives_restricted_vector_candidate():
    svc, _ = build_service()
    svc.index("salary confidential 90000", src("hr/salary.docx", "salary.docx", restricted="HR", department="HR"))
    svc.index("salary project budget 5000", src("projects/budget.txt", "budget.txt"))
    result = svc.retrieve(user(), RetrievalRequest("salary", top_k=1, candidate_k=10))
    assert result.hits
    assert result.hits[0].chunk.source.source_id == "projects/budget.txt"


def test_index_document_preserves_pdf_pages_and_excel_sheets():
    from backend.documents.models import Document, DocumentMetadata, DocumentSource
    access = AccessControlMetadata(tenant_id="eval")
    svc, _ = build_service()
    pdf_doc = Document(
        source=DocumentSource("report.pdf", "report.pdf", "pdf"),
        metadata=DocumentMetadata(10),
        text="[Page 1]\nAlpha revenue.\n\n[Page 2]\nBeta revenue.",
        locations=["Page 1", "Page 2"],
    )
    chunks = svc.index_document(pdf_doc, access)
    assert {c.source.page for c in chunks} == {1, 2}
    assert all(c.source.filename == "report.pdf" for c in chunks)

    xlsx_doc = Document(
        source=DocumentSource("sales.xlsx", "sales.xlsx", "xlsx"),
        metadata=DocumentMetadata(10),
        text="[Sheet January]\nRevenue 100.\n\n[Sheet February]\nRevenue 200.",
        locations=["Sheet January", "Sheet February"],
    )
    chunks = svc.index_document(xlsx_doc, access)
    assert {c.source.sheet for c in chunks} == {"January", "February"}


def test_source_terms_improve_ranking_without_increasing_top_k():
    svc, _ = build_service()
    svc.index("Customer renewal for Acme is due on 2026-11-30.", src("customers/renewal.pdf", "renewal.pdf"))
    svc.index("Acme renewal is due on 2026-11-30.", src("mail-1", "Acme renewal" , source_type="email"))
    result = svc.retrieve(user(), RetrievalRequest("Acme renewal PDF", top_k=1, candidate_k=10))
    assert len(result.hits) == 1
    assert result.hits[0].chunk.source.source_id == "customers/renewal.pdf"


def test_exact_month_term_ranks_matching_csv_above_generic_revenue_report():
    svc, _ = build_service()
    svc.index("Annual revenue summary Q1 120000 Q2 150000 Q3 130000.", src("reports/revenue-summary.pdf", "revenue-summary.pdf"))
    svc.index("Month Revenue January 100 February 200 March 300.", src("sales/monthly.csv", "monthly.csv"))
    result = svc.retrieve(user(), RetrievalRequest("March revenue", top_k=1, candidate_k=10))
    assert result.hits[0].chunk.source.source_id == "sales/monthly.csv"


def test_local_and_shared_file_sources_are_retrievable_with_source_identity():
    with TemporaryDirectory() as td:
        root = Path(td)
        (root / "local").mkdir()
        (root / "shared").mkdir()
        (root / "local" / "local.txt").write_text("Local handbook says office opens at 09:00.", encoding="utf-8")
        (root / "shared" / "shared.txt").write_text("Shared team calendar says planning is Friday.", encoding="utf-8")
        repo = LocalFileRepository(root)
        fs = FileService(repo, AuthorizationService(), AuditLogger(InMemoryAuditSink()))
        u = user()
        assert fs.get_metadata(u, "local/local.txt").name == "local.txt"
        assert fs.get_metadata(u, "shared/shared.txt").name == "shared.txt"

        svc, _ = build_service()
        svc.index(fs.read_file(u, "local/local.txt").decode(), src("local/local.txt", "local.txt"))
        svc.index(fs.read_file(u, "shared/shared.txt").decode(), src("shared/shared.txt", "shared.txt"))
        assert svc.retrieve(u, RetrievalRequest("office opens", 1, 5)).hits[0].chunk.source.source_id == "local/local.txt"
        assert svc.retrieve(u, RetrievalRequest("planning Friday", 1, 5)).hits[0].chunk.source.source_id == "shared/shared.txt"
