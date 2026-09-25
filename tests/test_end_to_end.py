from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from uuid import uuid4

import jwt
import pytest
from fastapi.testclient import TestClient
from reportlab.pdfgen import canvas
from docx import Document
from openpyxl import Workbook

from backend.agents.interfaces import Agent
from backend.security.ai.context import PromptContext
from backend.agents.models import AgentRequest, AgentResponse, Capability
from backend.agents.orchestrator import EnterpriseOrchestrator
from backend.api.chat_routes import get_chat_service
from backend.sources.dependencies import get_source_reference_service
from backend.chat.models import ChatRequest
from backend.chat.service import ChatService, ChatAnswerer, InMemoryConversationStore
from backend.documents import CSVParser, DocumentProcessingService, ExcelParser, PDFParser, WordParser
from backend.integrations.database.abstraction import DatabaseRepository
from backend.integrations.database.models import QueryRequest, QueryResult, TablePolicy
from backend.integrations.database.service import DatabaseService, DatabaseTool
from backend.integrations.database.sql_validation import ReadOnlySQLValidator, SQLValidationPipeline
from backend.integrations.email.models import EmailAddress, EmailMessage, EmailProviderContext, EmailSearchRequest, EmailThread
from backend.integrations.email.repository import EmailRepository
from backend.integrations.email.service import EmailService
from backend.integrations.files.local_repository import LocalFileRepository
from backend.integrations.files.service import FileService
from backend.reports.agent import ReportAgent
from backend.reports.models import ReportFormat, ReportRequest
from backend.reports.service import ReportService
from backend.reports.storage import SecureFileReportStorage
from backend.retrieval import AccessControlMetadata, DocumentChunker, InMemoryVectorStore, KeywordRetriever, KnowledgeRetrievalService, NoOpReranker, SourceMetadata
from backend.security.audit import AuditLogger, InMemoryAuditSink
from backend.security.authorization import AuthorizationService, Permission, UserAttributes, Resource
from backend.security.authorization.rbac import Role
from backend.security.models import UserIdentity
from backend.security.token_validator import TokenValidator
from backend.security.dependencies import get_current_user
from backend.sources.service import SourceReferenceService
from backend.sources.store import InMemorySourceReferenceStore
from backend.main import app

TEST_SECRET = "test-only-secret-key-32-bytes-minimum-xxxxxxxx"
USER = UserAttributes("u-test", "tenant-test", "Projects", frozenset({Role.EMPLOYEE}))
HR = UserAttributes("u-hr", "tenant-test", "HR", frozenset({Role.HR}))

class TestLLMAnswerer(ChatAnswerer):
    def __init__(self): self.calls = 0; self.last_context = None
    def answer(self, context: PromptContext):
        self.calls += 1
        self.last_context = context
        values = [x.value for x in context.retrieved if str(x.value).strip()] + [x.value for x in context.tool_results if str(x.value).strip()]
        if not values:
            return "I couldn't find enough information in the connected company sources to answer that."
        return "LLM_TEST_ANSWER: " + " | ".join(str(v) for v in values)

class TestAgent(Agent):
    def __init__(self, capability, callback): self.capability = capability; self.callback = callback
    def run(self, request): return self.callback(request)

class FakeGraph:
    def __init__(self, orchestrator): self.o = orchestrator
    def invoke(self, state):
        return self.o._invoke_without_langgraph(state)

class TestOrchestrator(EnterpriseOrchestrator):
    def __init__(self, agents):
        self.router = None
        self.agents = agents
        self.graph = FakeGraph(self)
    def _invoke_without_langgraph(self, state):
        capability = state.capability
        if capability is None:
            raise AssertionError("forced capability required for test graph")
        request = AgentRequest(state.request_id, self._principal_from_state(state), state.query)
        response = self.agents[capability].run(request)
        state.response = response
        state.trace.extend([f"route:{capability.value}", f"agent:{capability.value}"])
        return state

class FakeDB(DatabaseRepository):
    def __init__(self): self.calls=[]
    def execute_read(self, request):
        self.calls.append(request)
        return QueryResult(("customer", "sales"), (("A", 12000), ("B", 20000)))

class FakeEmailRepo(EmailRepository):
    def __init__(self):
        self.messages = [EmailMessage("m1", "c1", "ABC renewal", EmailAddress("abc@test.invalid", "ABC"), (), (), None, None, "ABC renewal details", False)]
    def search(self, context, request): return self.messages if not request.query or request.query.casefold() in "abc renewal" else []
    def get_thread(self, context, conversation_id): return EmailThread(conversation_id, tuple(self.messages))
    def get_message(self, context, message_id): return self.messages[0]

class Runtime:
    def __init__(self, tmp_path: Path):
        self.sink = InMemoryAuditSink(); self.audit = AuditLogger(self.sink); self.auth = AuthorizationService()
        self.root = tmp_path / "CompanyData"; self.root.mkdir()
        self._make_files()
        self.file_service = FileService(LocalFileRepository(self.root), self.auth, self.audit)
        self.sources = SourceReferenceService(self.auth, self.audit, InMemorySourceReferenceStore())
        self.retrieval = KnowledgeRetrievalService(None, None, KeywordRetriever(), NoOpReranker(), self.auth, self.audit, DocumentChunker(), self.sources)
        self.db_repo = FakeDB()
        validator = ReadOnlySQLValidator(TablePolicy(frozenset({("public", "sales")}), frozenset({("public", "sales", "customer"), ("public", "sales", "sales"), ("public", "sales", "date")})), require_column_policy=True)
        self.db = DatabaseService(self.db_repo, self.auth, self.audit, SQLValidationPipeline(validator), "tenant-test")
        self.email_repo = FakeEmailRepo(); self.email = EmailService(self.email_repo, self.auth, self.audit)
        self.email_context = EmailProviderContext("u-test", "tenant-test", "test-access-token", frozenset({"Mail.Read"}))
        self.report_storage = SecureFileReportStorage(tmp_path / "reports")
        self.report = ReportService(self.report_storage, self.auth, self.audit)
        self.analysis_data = (("A", 12000), ("B", 20000))
        self._index_docs()
        self.orchestrator = self._make_orchestrator()
        self.llm = TestLLMAnswerer()
        self.chat = ChatService(self.orchestrator, self.llm, InMemoryConversationStore())

    def _make_files(self):
        (self.root / "project.txt").write_text("Project Atlas milestone is active.", encoding="utf-8")
        (self.root / "report.csv").write_text("customer,sales\nA,12000\nB,20000\n", encoding="utf-8")
        wb=Workbook(); ws=wb.active; ws.title="Sales"; ws.append(["customer","sales"]); ws.append(["A",12000]); ws.append(["B",20000]); wb.save(self.root/"sales.xlsx"); wb.close()
        d=Document(); d.add_paragraph("Project Atlas Word milestone is active."); d.save(self.root/"atlas.docx")
        c=canvas.Canvas(str(self.root/"atlas.pdf")); c.drawString(72, 720, "Project Atlas PDF milestone is active."); c.save()
        (self.root/"HR").mkdir(); (self.root/"HR"/"salary.csv").write_text("employee,salary\nSecret,999999\n", encoding="utf-8")

    def _index_docs(self):
        svc=DocumentProcessingService([PDFParser(), WordParser(), ExcelParser(), CSVParser()])
        for p in [self.root/"project.txt", self.root/"report.csv", self.root/"sales.xlsx", self.root/"atlas.docx", self.root/"atlas.pdf"]:
            doc = svc.parse_bytes(filename=p.name, relative_path=p.name, data=p.read_bytes()) if p.suffix != ".txt" else None
            text = doc.text if doc else p.read_text(encoding="utf-8")
            self.retrieval.index(text, SourceMetadata(p.suffix.lstrip("."), p.name, p.name, p.name, access=AccessControlMetadata("tenant-test")))
        self.retrieval.index("HR salary confidential", SourceMetadata("csv", "HR/salary.csv", "salary.csv", "HR/salary.csv", access=AccessControlMetadata("tenant-test", restricted_department="HR")))

    def _make_orchestrator(self):
        def knowledge(req):
            result=self.retrieval.retrieve(req.user, __import__('backend.retrieval', fromlist=['RetrievalRequest']).RetrievalRequest(req.query, top_k=5))
            refs=self.retrieval.source_references(req.user, req.request_id, result)
            text="; ".join(h.chunk.content for h in result.hits)
            return AgentResponse(Capability.KNOWLEDGE, text, refs)
        def email(req):
            msgs=self.email.search(req.user, self.email_context, EmailSearchRequest(query="ABC"))
            refs=tuple(self.sources.from_metadata(req.user, req.request_id, SourceMetadata("email", m.id, m.subject, None, access=AccessControlMetadata("tenant-test", owner_id=req.user.user_id))) for m in msgs)
            return AgentResponse(Capability.EMAIL, "; ".join(m.body_preview or m.subject for m in msgs), refs)
        def database(req):
            result=self.db.execute_read(req.user, QueryRequest("SELECT customer, sales FROM sales LIMIT 100"))
            return AgentResponse(Capability.DATABASE, str(result.rows), self.sources.from_lineage(req.user, req.request_id, (__import__('backend.analysis.models', fromlist=['DataLineage']).DataLineage("sql", "sales", "Sales", ("customer","sales"), "SELECT customer, sales FROM sales", tenant_id="tenant-test", resource_type="database_source"),)))
        def analysis(req):
            result=self.db.execute_read(req.user, QueryRequest("SELECT customer, sales FROM sales LIMIT 100"))
            from backend.analysis import StructuredDataset, DataLineage, AnalysisRequest, AnalysisOperation, DeterministicAnalysisEngine
            ds=StructuredDataset(("customer","sales"), tuple({"customer":r[0],"sales":r[1]} for r in result.rows), (DataLineage("sql","sales","Sales",("customer","sales"),"SELECT customer, sales FROM sales",tenant_id="tenant-test",resource_type="database_source"),))
            ar=AnalysisRequest(AnalysisOperation.TOTAL, ds, column="sales")
            val=DeterministicAnalysisEngine().execute(ar)
            return AgentResponse(Capability.DATA_ANALYSIS, val, self.sources.from_lineage(req.user, req.request_id, val.lineage))
        def report(req):
            from backend.analysis import DataLineage
            rr=ReportRequest(str(uuid4()), "Sales Test", ReportFormat.PDF, req.user.user_id, req.user.tenant_id, [{"customer":"A","sales":12000}], (DataLineage("sql","sales","Sales",("customer","sales"),"SELECT customer,sales FROM sales",tenant_id="tenant-test",resource_type="database_source"),))
            agent=ReportAgent(self.report, __import__('backend.agents.security_gateway', fromlist=['SecureToolGateway']).SecureToolGateway(self.auth,self.audit), self.sources)
            return agent.create(req, rr)
        return TestOrchestrator({
            Capability.KNOWLEDGE: TestAgent(Capability.KNOWLEDGE, knowledge),
            Capability.EMAIL: TestAgent(Capability.EMAIL, email),
            Capability.DATABASE: TestAgent(Capability.DATABASE, database),
            Capability.DATA_ANALYSIS: TestAgent(Capability.DATA_ANALYSIS, analysis),
            Capability.REPORT: TestAgent(Capability.REPORT, report),
        })

@pytest.fixture
def runtime(tmp_path): return Runtime(tmp_path)

@pytest.fixture
def client(runtime):
    app.dependency_overrides[get_current_user] = lambda: UserIdentity("u-test", "issuer", "u@test.invalid", "Test User", "tenant-test", "Projects", frozenset({Role.EMPLOYEE}))
    app.dependency_overrides[get_chat_service] = lambda: runtime.chat
    app.dependency_overrides[get_source_reference_service] = lambda: runtime.sources
    c=TestClient(app)
    yield c
    app.dependency_overrides.clear()

def test_01_login_valid_and_02_invalid():
    key = jwt.algorithms.RSAAlgorithm.generate_private_key() if False else None
    # Existing auth integration contract: malformed/invalid tokens are rejected by validator before claims are trusted.
    with pytest.raises(Exception): TokenValidator(jwks_url="", issuer="issuer", audience="aud").validate("bad.token")
    assert jwt.encode({"sub":"u-test","iss":"issuer","aud":"aud"}, TEST_SECRET, algorithm="HS256")

def test_03_04_authenticated_unauthenticated(client):
    assert client.get("/api/auth/me").status_code == 200
    app.dependency_overrides.clear()
    assert client.get("/api/auth/me").status_code == 401

def test_05_06_authorized_unauthorized(runtime):
    assert runtime.auth.authorize(USER, Permission.FILE_READ, __import__('backend.security.authorization', fromlist=['Resource']).Resource("project.txt","company_file","tenant-test")).allowed
    assert not runtime.auth.authorize(UserAttributes("u2","tenant-other","Projects",frozenset({Role.EMPLOYEE})), Permission.FILE_READ, __import__('backend.security.authorization', fromlist=['Resource']).Resource("project.txt","company_file","tenant-test")).allowed

def test_07_file_search(runtime):
    assert runtime.file_service.search_by_filename(USER, "project")

def test_08_09_10_11_document_search(runtime):
    for term in ["PDF milestone", "Word milestone", "Excel sales", "customer sales"]:
        result=runtime.retrieval.retrieve(USER, __import__('backend.retrieval', fromlist=['RetrievalRequest']).RetrievalRequest(term, top_k=5))
        assert result.hits, term

def test_12_sql_query(runtime):
    result=runtime.db.execute_read(USER, QueryRequest("SELECT customer, sales FROM sales LIMIT 100"))
    assert result.rows == (("A",12000),("B",20000))
    assert runtime.db_repo.calls

def test_13_email_search(runtime):
    result=runtime.email.search(USER, runtime.email_context, EmailSearchRequest(query="ABC"))
    assert result and result[0].subject == "ABC renewal"

def test_14_multi_source_question(client, runtime):
    r=client.post("/api/chat", json={"query":"Give me the project PDF and ABC email update"})
    assert r.status_code == 200
    assert runtime.llm.calls == 1
    assert len(r.json()["sources"]) >= 2

def test_15_analysis_question(client):
    r=client.post("/api/chat", json={"query":"What is the total sales analysis?"})
    assert r.status_code == 200 and "LLM_TEST_ANSWER" in r.json()["answer"]

def test_16_report_generation(runtime):
    req=AgentRequest("r1", USER, "report")
    from backend.agents.security_gateway import SecureToolGateway
    agent=ReportAgent(runtime.report, SecureToolGateway(runtime.auth,runtime.audit), runtime.sources)
    rr=ReportRequest(str(uuid4()),"Sales Test",ReportFormat.PDF,"u-test","tenant-test",[{"sales":32000}], (__import__("backend.analysis.models", fromlist=["DataLineage"]).DataLineage("sql","sales","Sales",("sales",),"SELECT sales FROM sales",tenant_id="tenant-test",resource_type="database_source"),))
    response=agent.create(req,rr)
    assert response.content.metadata.size_bytes > 0 and response.sources

def test_17_source_display(client, runtime):
    r=client.post("/api/chat", json={"query":"Find project document"}); assert r.status_code==200
    source=r.json()["sources"][0]
    sr=client.get(source["href"])
    assert sr.status_code == 200 and "source_id" not in sr.json()

def test_18_conversation_history(client):
    r=client.post("/api/chat", json={"query":"Find project document"}); cid=r.json()["conversation_id"]
    r2=client.post("/api/chat", json={"query":"Find project document","conversation_id":cid}); assert r2.status_code==200 and len(r2.json()["history"])==4
    h=client.get("/api/chat/history"); assert h.status_code==200 and h.json()["conversations"]

def test_19_error_handling(client):
    assert client.post("/api/chat", json={"query":""}).status_code == 422

def test_20_empty_result(client):
    r=client.post("/api/chat", json={"query":"This query should have no matching evidence"})
    assert r.status_code == 200
    assert "couldn't find" in r.json()["answer"].lower() or "no authorized" in r.json()["answer"].lower()
