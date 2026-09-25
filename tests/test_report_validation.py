from __future__ import annotations

import re
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from backend.analysis.models import DataLineage
from backend.reports.models import ReportFormat, ReportRequest
from backend.reports.service import ReportAuthorizationError, ReportService
from backend.reports.storage import SecureFileReportStorage
from backend.security.audit import AuditLogger, InMemoryAuditSink
from backend.security.authorization import AuthorizationService, UserAttributes
from backend.security.authorization.rbac import Role


def user(role=Role.EMPLOYEE, user_id="u1", tenant="t1", department="Projects"):
    return UserAttributes(user_id, tenant, department, frozenset({role}))


def lineage(source_id="sales", label="Sales Data", *, tenant="t1", department=None, restricted=None):
    return DataLineage(
        "sql", source_id, label, ("month", "sales"),
        "SELECT month, sales FROM sales", (),
        tenant_id=tenant, department=department, restricted_department=restricted,
        resource_type="database_source",
    )


def req(fmt, data, *, owner="u1", tenant="t1", title="Sales / Summary: Q3?", lin=None):
    return ReportRequest(str(uuid4()), title, fmt, owner, tenant, data,
                         (lin or lineage(tenant=tenant),), "Calculated from authorized data")


def service(tmp_path):
    sink = InMemoryAuditSink()
    return ReportService(SecureFileReportStorage(tmp_path / "reports"), AuthorizationService(), AuditLogger(sink)), sink


@pytest.mark.parametrize("fmt, suffix", [
    (ReportFormat.EXCEL, ".xlsx"),
    (ReportFormat.PDF, ".pdf"),
    (ReportFormat.WORD, ".docx"),
    (ReportFormat.POWERPOINT, ".pptx"),
])
def test_all_supported_formats_are_real_and_contain_correct_data(tmp_path, fmt, suffix):
    svc, _ = service(tmp_path)
    data = [{"month": "Jan", "sales": 10000}, {"month": "Feb", "sales": 25000}]
    artifact = svc.create_report(user(), req(fmt, data))
    path = svc._storage.open_path(artifact.metadata)
    assert path.suffix == suffix and path.stat().st_size > 0

    if fmt == ReportFormat.EXCEL:
        wb = load_workbook(path, read_only=True, data_only=True)
        values = [[c.value for c in row] for row in wb["Report"].iter_rows()]
        assert "sales" in values[3]
        assert 10000 in [v for row in values for v in row]
        assert 25000 in [v for row in values for v in row]
        assert wb["Sources"]["A2"].value == "sql"
    elif fmt == ReportFormat.PDF:
        text = "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
        assert "Jan" in text and "25000" in text
        assert "Sales Data" in text
    elif fmt == ReportFormat.WORD:
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        table_text = "\n".join(c.text for row in doc.tables[0].rows for c in row.cells)
        assert "Sales Data" in text and "Jan" in table_text and "25000" in table_text
    else:
        prs = Presentation(path)
        cell_text = "\n".join(
            cell.text for slide in prs.slides for shape in slide.shapes
            if getattr(shape, "has_table", False) for row in shape.table.rows for cell in row.cells
        )
        all_text = "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text_frame"))
        assert "Jan" in cell_text and "25000" in cell_text and "Sales Data" in all_text


def test_report_authorizes_every_source_before_generation(tmp_path):
    svc, _ = service(tmp_path)
    finance = lineage("finance", "Finance Data", department="Finance", restricted="Finance")
    with pytest.raises(ReportAuthorizationError):
        svc.create_report(user(), req(ReportFormat.PDF, [{"salary": 999999}], lin=finance))
    assert not list((tmp_path / "reports").glob("*.pdf"))


def test_allowed_user_can_generate_restricted_department_data(tmp_path):
    svc, _ = service(tmp_path)
    finance = lineage("finance", "Finance Data", department="Finance", restricted="Finance")
    artifact = svc.create_report(user(Role.FINANCE, department="Finance"), req(
        ReportFormat.EXCEL, [{"month": "Jan", "sales": 50000}], owner="u1", lin=finance))
    assert artifact.metadata.size_bytes > 0


def test_user_b_different_allowed_data_cannot_reuse_user_a_report(tmp_path):
    svc, _ = service(tmp_path)
    a = user(user_id="u1")
    b = user(user_id="u2")
    artifact = svc.create_report(a, req(ReportFormat.PDF, [{"sales": 10}], owner="u1"))
    with pytest.raises(ReportAuthorizationError):
        svc.issue_temporary_download_url(b, artifact.metadata.report_id)


def test_cross_tenant_lineage_is_rejected(tmp_path):
    svc, _ = service(tmp_path)
    with pytest.raises(ReportAuthorizationError):
        svc.create_report(user(), req(ReportFormat.PDF, [{"sales": 10}], lin=lineage(tenant="other")))


def test_missing_lineage_is_rejected_before_generation(tmp_path):
    svc, _ = service(tmp_path)
    bad = ReportRequest(str(uuid4()), "No Provenance", ReportFormat.EXCEL, "u1", "t1", [{"secret": "x"}], ())
    with pytest.raises(ReportAuthorizationError):
        svc.create_report(user(), bad)
    assert not list((tmp_path / "reports").glob("*.xlsx"))


def test_filename_is_safe_and_extension_is_fixed(tmp_path):
    svc, _ = service(tmp_path)
    title = 'Q3/Payroll\\2026\r\n"<secret>"'
    artifact = svc.create_report(user(), req(ReportFormat.PDF, [{"sales": 1}], title=title))
    from backend.api.report_routes import _safe_download_filename
    filename = _safe_download_filename(artifact.metadata.title, artifact.metadata.format.value)
    assert not re.search(r'[\\/:*?"<>|\r\n]', filename)
    assert filename.endswith(".pdf")


def test_empty_dataset_is_valid_in_all_formats(tmp_path):
    svc, _ = service(tmp_path)
    for fmt in ReportFormat:
        artifact = svc.create_report(user(), req(fmt, []))
        path = svc._storage.open_path(artifact.metadata)
        assert path.stat().st_size > 0
        if fmt == ReportFormat.PDF:
            assert PdfReader(str(path)).pages
        elif fmt == ReportFormat.WORD:
            assert Document(path).paragraphs
        elif fmt == ReportFormat.POWERPOINT:
            assert len(Presentation(path).slides) >= 2
        elif fmt == ReportFormat.TEXT:
            assert len(path.read_text(encoding="utf-8")) > 0
        else:
            assert load_workbook(path, read_only=True)["Report"].max_row >= 4


def test_large_dataset_is_split_or_stored_without_loss(tmp_path):
    svc, _ = service(tmp_path)
    data = [{"id": i, "value": i * 2} for i in range(1000)]
    for fmt in (ReportFormat.EXCEL, ReportFormat.POWERPOINT, ReportFormat.WORD):
        artifact = svc.create_report(user(), req(fmt, data))
        path = svc._storage.open_path(artifact.metadata)
        assert path.stat().st_size > 0
        if fmt == ReportFormat.EXCEL:
            wb = load_workbook(path, read_only=True, data_only=True)
            assert wb["Report"].max_row == 1004
        elif fmt == ReportFormat.POWERPOINT:
            prs = Presentation(path)
            assert len(prs.slides) > 10


def test_excel_formula_injection_remains_neutralized(tmp_path):
    svc, _ = service(tmp_path)
    artifact = svc.create_report(user(), req(ReportFormat.EXCEL, [{"name": "=HYPERLINK(\"https://evil.example\")"}]))
    wb = load_workbook(svc._storage.open_path(artifact.metadata), read_only=True, data_only=False)
    assert str(wb["Report"]["A5"].value).startswith("'")


def test_download_is_permission_checked_and_one_time(tmp_path):
    svc, _ = service(tmp_path)
    artifact = svc.create_report(user(), req(ReportFormat.PDF, [{"sales": 1}]))
    token_url = svc.issue_temporary_download_url(user(), artifact.metadata.report_id)
    token = token_url.split("token=", 1)[1]
    path, _ = svc.open_download(user(), artifact.metadata.report_id, token)
    assert path.exists()
    with pytest.raises(Exception):
        svc.open_download(user(), artifact.metadata.report_id, token)


def test_calculated_values_are_preserved_exactly(tmp_path):
    svc, _ = service(tmp_path)
    data = [{"total_sales": 35000, "average_sales": 17500, "change_pct": 40.0}]
    artifact = svc.create_report(user(), req(ReportFormat.EXCEL, data))
    wb = load_workbook(svc._storage.open_path(artifact.metadata), read_only=True, data_only=True)
    values = [c.value for row in wb["Report"].iter_rows(min_row=5) for c in row]
    assert values == [35000, 17500, 40.0]


def test_special_characters_in_content_are_preserved_safely(tmp_path):
    svc, _ = service(tmp_path)
    data = [{"name": "Zoë & R&D <North>", "note": "Line 1 / Line 2 — 日本語"}]
    for fmt in ReportFormat:
        artifact = svc.create_report(user(), req(fmt, data, title="Q3 — R&D <Summary>"))
        path = svc._storage.open_path(artifact.metadata)
        assert path.stat().st_size > 0
        if fmt == ReportFormat.EXCEL:
            wb = load_workbook(path, read_only=True, data_only=True)
            vals = [c.value for row in wb["Report"].iter_rows() for c in row]
            assert "Zoë & R&D <North>" in vals
        elif fmt == ReportFormat.PDF:
            text = "\n".join(p.extract_text() or "" for p in PdfReader(str(path)).pages)
            assert "Q3" in text and "R&D" in text
        elif fmt == ReportFormat.WORD:
            doc = Document(path)
            assert "Zoë & R&D <North>" in "\n".join(c.text for row in doc.tables[0].rows for c in row.cells)
        elif fmt == ReportFormat.TEXT:
            assert "Zoë & R&D <North>" in path.read_text(encoding="utf-8")
        else:
            prs = Presentation(path)
            cells = [cell.text for slide in prs.slides for shape in slide.shapes if getattr(shape, "has_table", False) for row in shape.table.rows for cell in row.cells]
            assert "Zoë & R&D <North>" in cells


def test_mixed_authorized_and_unauthorized_sources_are_rejected_atomically(tmp_path):
    svc, _ = service(tmp_path)
    finance = lineage("finance", "Finance Data", department="Finance", restricted="Finance")
    mixed = (lineage(), finance)
    bad = ReportRequest(str(uuid4()), "Mixed", ReportFormat.PDF, "u1", "t1", [{"sales": 10, "salary": 999999}], mixed)
    with pytest.raises(ReportAuthorizationError):
        svc.create_report(user(), bad)
    assert not list((tmp_path / "reports").glob("*.pdf"))


def test_invalid_runtime_format_fails_safely(tmp_path):
    svc, _ = service(tmp_path)
    bad = ReportRequest(str(uuid4()), "Bad", "html", "u1", "t1", [{"x": 1}], (lineage(),))
    with pytest.raises(ValueError, match="Unsupported report format"):
        svc.create_report(user(), bad)


def test_pdf_large_dataset_is_readable(tmp_path):
    svc, _ = service(tmp_path)
    data = [{"id": i, "value": i * 3} for i in range(500)]
    artifact = svc.create_report(user(), req(ReportFormat.PDF, data))
    reader = PdfReader(str(svc._storage.open_path(artifact.metadata)))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "499" in text and "1497" in text
